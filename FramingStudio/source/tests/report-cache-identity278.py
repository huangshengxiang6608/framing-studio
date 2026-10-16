"""Read-only regression for original B-copy cache routing and printed errors.

Usage: python report-cache-identity278.py VERIFIED_B_RUN
"""
from pathlib import Path
import json
import math
import re
import sys
import openpyxl
from pypdf import PdfReader

root = Path(sys.argv[1])
result = json.loads((root / "result.json").read_text(encoding="utf-8-sig"))
job = json.loads((root / "job.json").read_text(encoding="utf-8-sig"))
assert result["ok"] and not result["differences"] and job["section"] == "B"
fields = ["C40", "D40", "C41", "D41", "C42", "D42", "C43", "D43", "C44", "D44", "C45", "D45", "C46", "D46", "C47", "D47", "E58", "G58", "H58", "E76", "G76", "H76", "T8", "T9", "T10", "T11", "C110", "D110", "C111", "D111", "C112", "D112", "C113", "D113", "C114", "D114", "C115", "D115", "C116", "D116", "C117", "D117"]
checked = 0
displaced = 0
parts = [(root / f"column-part-{i+1}", p) for i, p in enumerate(job["columnParts274"])] if "columnParts274" in job else [(root, job)]
for directory, part in parts:
    book = openpyxl.load_workbook(directory / "RC-filled.xlsm", data_only=True)
    for member in part["members"]:
        if member["kind"] not in ("MB", "SB") or not member.get("adopt254"):
            continue
        sheet = book[f"_{member['kind']} Analysis"]
        rows = [r for r in range(6, 26) if sheet.cell(r, 200).value == member["id"]]
        assert len(rows) == 1, (member["id"], rows)
        row = rows[0]
        displaced += row != member["row"] - (68 if member["kind"] == "MB" else 43)
        for address, expected in member["adopt254"]["steel"].items():
            actual = sheet.cell(row, 150 + fields.index(address)).value
            if isinstance(actual, str) and actual.startswith("="):
                continue  # Native input formula is intentionally cached as a formula.
            if expected in (None, ""):
                assert actual in (None, ""), (member["id"], address, actual)
            else:
                assert isinstance(actual, (int, float)) and math.isclose(actual, expected), (member["id"], address, expected, actual)
            checked += 1
pages = 0
for f in result["files"]:
    for page in PdfReader(root / f["file"]).pages:
        text = page.extract_text()
        assert not re.search(r"#(?:VALUE!|REF!|DIV/0!|NUM!|NAME\?|N/A)", text), f["file"]
        pages += 1
assert checked and displaced, "Use a model containing non-positional beam cache entries"
print(json.dumps({"cachedSteelInputs": checked, "nonPositionalMembers": displaced, "pdfPagesWithoutCellErrors": pages}))
