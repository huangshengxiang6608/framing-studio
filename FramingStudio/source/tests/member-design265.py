"""Read-only validation of per-member native Excel output: python member-design265.py RUN_DIR."""
import json,re,sys,math
from pathlib import Path
import openpyxl

d=Path(sys.argv[1]);result=json.loads((d/'result.json').read_text(encoding='utf-8-sig'))
assert result['ok'] and not result['differences']
job=json.loads((d/'job.json').read_text(encoding='utf-8-sig'))
formula=openpyxl.load_workbook(d/result['workbook'],read_only=False,data_only=False)
values=openpyxl.load_workbook(d/result['workbook'],read_only=False,data_only=True)
counts={};checked=0
for member in job['members']:
 if member.get('selected') is False:continue
 name=re.sub(r'[\\/:?*\[\]]','',member.get('designName265',member['id'])).strip().strip("'")[:31]
 sheet=formula[name];source=formula[member['sheet']]
 assert not sheet.protection.sheet and sheet.sheet_state=='visible'
 assert '_RCDesign265' in sheet.defined_names and sheet.defined_names['_RCDesign265'].attr_text=='TRUE'
 assert formula.sheetnames.index(name)>formula.sheetnames.index(source.title)
 previous=counts.get(source.title,source.title);assert formula.sheetnames.index(name)==formula.sheetnames.index(previous)+1;counts[source.title]=name
 source_formulas={c.coordinate for row in source for c in row if c.data_type=='f'}
 copied_formulas={c.coordinate for row in sheet for c in row if c.data_type=='f'}
 # Native CB report generation restores G13's source formula after capture;
 # its adopted eccentricity was a checked input at capture (including zero).
 if member['kind']=='CB' and 'G13' in source_formulas-copied_formulas:
  assert values[name]['G13'].value==(member.get('adopt254') or member)['inputs']['G13']
  source_formulas.remove('G13')
 assert source_formulas==copied_formulas,(name,'lost original calculation formulas')
 for row in sheet:
  for cell in row:
   if cell.data_type=='f':
    for other in formula.sheetnames:
     if other.startswith('Section ') or other in ('Input','_Input Engine'):
      assert "'"+other+"'!" not in cell.value,(name,cell.coordinate,'shared input reference')
 data=member.get('adopt254') or member
 for field in ('inputs','expected','steel'):
  for address,expected in (data.get(field) or {}).items():
   actual=values[name][address].value
   if address=='G13' and actual is None and expected==0:actual=0
   if isinstance(expected,(float,int)) and not isinstance(expected,bool):assert actual is not None and math.isclose(expected,actual,rel_tol=1e-7,abs_tol=1e-7),(name,address,expected,actual)
   else:assert ('' if actual is None else actual)==('' if expected is None else expected),(name,address,expected,actual)
   checked+=1
print(json.dumps({'members':len(job['members']),'formulaCoordinatesPreserved':True,'inputResultCells':checked,'independentInputReferences':True}))
